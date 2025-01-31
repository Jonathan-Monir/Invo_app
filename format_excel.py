import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import PatternFill, Font, Border, Side, Alignment
import os

class FormatExcel:
    def __init__(self, df, output_file_path="formatted_output.xlsx"):
        self.df = df
        self.output_file_path = output_file_path
        self.workbook = Workbook()
        self.worksheet = self.workbook.active
        self.write_dataframe_to_sheet()
        self.change_colors()
        self.change_font()
        self.set_column_width()
        self.save_output_file()

    def write_dataframe_to_sheet(self):
        # Define border style
        border = Border(left=Side(style='thin'), 
                        right=Side(style='thin'), 
                        top=Side(style='thin'), 
                        bottom=Side(style='thin'))

        # Write headers with borders and alignment
        for col_idx, col_name in enumerate(self.df.columns, start=1):
            cell = self.worksheet.cell(row=1, column=col_idx, value=col_name)
            cell.border = border  # Apply border to headers
            cell.alignment = Alignment(horizontal='center', vertical='center')

        # Write DataFrame to worksheet
        for row_idx, row in enumerate(self.df.itertuples(index=False), start=2):
            fill = PatternFill(start_color='f0f0f0', end_color='f0f0f0', fill_type='solid') if row_idx % 2 == 0 else PatternFill()
            for col_idx, value in enumerate(row, start=1):
                cell = self.worksheet.cell(row=row_idx, column=col_idx, value=value)
                cell.fill = fill
                cell.alignment = Alignment(horizontal='center', vertical='center')
                cell.border = border  # Apply border to data cells

    
    def save_output_file(self):
        if os.path.exists(self.output_file_path):
            os.remove(self.output_file_path)
        self.workbook.save(self.output_file_path)
        print(f"Saved at {self.output_file_path}")

    def change_colors(self):
        fill = PatternFill(start_color='ff4300', end_color='ff4300', fill_type='solid')
        fill_first_row = PatternFill(start_color='ffff00', end_color='ffff00', fill_type='solid')

        # Find "Difference" column index
        amount_col_index = None
        for col_idx, col_name in enumerate(self.df.columns, start=1):
            if col_name == "Difference":
                amount_col_index = col_idx
                break

        if amount_col_index is not None:
            for row_idx in range(2, len(self.df) + 2):
                amount_cell = self.worksheet.cell(row=row_idx, column=amount_col_index)
                if amount_cell.value != 0:
                    amount_cell.fill = fill

        # Change color of first row (headers)
        for col_idx in range(1, len(self.df.columns) + 1):
            cell = self.worksheet.cell(row=1, column=col_idx)
            cell.fill = fill_first_row

    def change_font(self):
        font = Font(name='Calibri', size=14, bold=True, italic=False, color='000080')
        for col_idx in range(1, len(self.df.columns) + 1):
            cell = self.worksheet.cell(row=1, column=col_idx)
            cell.font = font

    def set_column_width(self):
        for col_idx, col_name in enumerate(self.df.columns, start=1):
            max_length = min(max(self.df[col_name].astype(str).apply(len).max(), len(col_name)),30) + 3
            self.worksheet.column_dimensions[self.worksheet.cell(row=1, column=col_idx).column_letter].width = max_length

if __name__ == "__main__":
    # Example usage
    data = {
        "Name": ["Alice", "Bob", "Charlie"],
        "Age": [25, 30, 35],
        "Difference": [0, 10, -5]
    }
    df = pd.DataFrame(data)
    formatter = FormatExcel(df, "formatted_output.xlsx")
    print("Excel file formatted successfully!")
